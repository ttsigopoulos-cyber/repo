"""Regenerates lib/questions.ts from the ICP questionnaire workbook.

Usage (from the project root):
    pip install openpyxl
    python scripts/extract_questions.py path/to/260930TT_ESCP_ICP_SCB_CareHousesGermany_Questionnaire.xlsx
"""
import json, re, sys
import openpyxl

src = sys.argv[1]
wb = openpyxl.load_workbook(src, data_only=True)
ws = wb.worksheets[0]  # "QUESTIONAIRE 2026CW39"
hdr = [c.value for c in ws[7]]
qs = []
for r in ws.iter_rows(min_row=8, values_only=True):
    d = dict(zip(hdr, r))
    if not d.get("Question"):
        continue
    code = d["Question"]; part = code[0]; blk = int(code[1:3])
    tier = int(re.search(r"Tier (\d)", d["Interview Priority"]).group(1))
    omit = str(d["Omission Allowed"])
    text = d["Improved Question German"].strip()
    variant = None
    m = re.search(r"\s*\[Für Heimleitung: … (.*?)\]\s*$", text)
    if m:  # A04.1 carries a Heimleitung variant in square brackets
        base = text[: m.start()].strip()
        end = base.index("stand?") + len("stand?")
        first, tail = base[:end], base[end:]
        cut = first.index("zwischen Ihnen und Ihrer eigentlichen Aufgabe stand?")
        variant = first[:cut] + m.group(1).rstrip(".") + "?" + tail
        text = base
    if part in ("A", "B", "D"):
        aud = ["pflegekraft", "leitung"]
    else:  # Part C: blocks 1–4 management, 5 care staff, 6–9 relatives (see sheet "Legend")
        aud = ["leitung"] if blk <= 4 else (["pflegekraft"] if blk == 5 else ["angehoerige"])
    qs.append(dict(seq=d["SQ"], code=code, part=part, block=blk, blockTitle=d["Block"], tier=tier,
                   neverSkip=omit.startswith("No – never skip"),
                   familyStopOnDistress=omit.startswith("Yes – and must be omitted"),
                   reactiveOnly=part == "D", audiences=aud, risk=d["Legal Risk"],
                   text=text, textLeitung=variant))

with open("lib/questions.ts", encoding="utf-8") as f:
    header = f.read().split("export const QUESTIONS")[0]
with open("lib/questions.ts", "w", encoding="utf-8") as f:
    f.write(header + "export const QUESTIONS: Question[] = " + json.dumps(qs, ensure_ascii=False, indent=2) + ";\n")
print(f"{len(qs)} questions written to lib/questions.ts")
